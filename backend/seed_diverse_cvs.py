import io
import os
import sys

# Configure UTF-8 stdout for Windows
if sys.stdout and hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

# Ensure backend path is on sys.path
backend_dir = os.path.dirname(os.path.abspath(__file__))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

import docx
import fitz  # PyMuPDF
from database.db import (
    application_exists,
    create_application,
    create_candidate,
    get_candidates,
    get_connection,
)
from pypdf import PdfReader
from scripts.rebuild_rag_index import rebuild_index

UPLOAD_DIR = os.path.join(backend_dir, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)
ARIAL_FONT = "C:/Windows/Fonts/arial.ttf"


def create_pdf_cv(filename, name, title, email, phone, sections):
    """Generate a clean styled PDF CV using PyMuPDF and Arial font."""
    file_path = os.path.join(UPLOAD_DIR, filename)
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4 size

    # Register font if exists
    if os.path.exists(ARIAL_FONT):
        page.insert_font(fontname="Arial", fontfile=ARIAL_FONT)
        font = "Arial"
    else:
        font = "helv"

    # Header background banner
    rect_header = fitz.Rect(0, 0, 595, 110)
    page.draw_rect(rect_header, color=None, fill=(0.12, 0.22, 0.42))

    page.insert_text((40, 42), name, fontsize=19, fontname=font, color=(1, 1, 1))
    page.insert_text((40, 66), title.upper(), fontsize=11, fontname=font, color=(0.82, 0.89, 1))
    page.insert_text(
        (40, 88),
        f"Email: {email}   |   SĐT: {phone}",
        fontsize=9.5,
        fontname=font,
        color=(0.92, 0.92, 0.92),
    )

    y = 140
    for sec_title, sec_content in sections:
        if y > 760:
            break
        page.insert_text((40, y), sec_title, fontsize=12, fontname=font, color=(0.12, 0.22, 0.42))
        page.draw_line((40, y + 4), (555, y + 4), color=(0.8, 0.85, 0.9), width=1)
        y += 20

        # Draw content lines or rect textbox
        rect_content = fitz.Rect(40, y, 555, y + 100)
        # Using insert_textbox for clean multiline wrapping
        rc = page.insert_textbox(
            rect_content,
            sec_content,
            fontsize=9.5,
            fontname=font,
            color=(0.18, 0.18, 0.18),
            lineheight=1.35,
        )
        # advance y based on lines
        lines_count = sec_content.count("\n") + int(len(sec_content) / 80) + 1
        y += max(lines_count * 15, 35) + 15

    doc.save(file_path)
    doc.close()
    return file_path


def create_docx_cv(filename, name, title, email, phone, sections):
    """Generate a clean styled DOCX CV using python-docx."""
    file_path = os.path.join(UPLOAD_DIR, filename)
    doc = docx.Document()

    # Title & Header
    h1 = doc.add_heading(name, 0)
    p_sub = doc.add_paragraph()
    p_sub.add_run(f"{title.upper()}\n").bold = True
    p_sub.add_run(f"Email: {email} | SĐT: {phone}").italic = True

    # Sections
    for sec_title, sec_content in sections:
        doc.add_heading(sec_title, level=1)
        for line in sec_content.strip().split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("- ") or line.startswith("• ") or line.startswith("+ "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)

    doc.save(file_path)
    return file_path


def extract_text(file_path):
    """Extract text from file using pypdf or python-docx."""
    ext = file_path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if ext == "docx":
        d = docx.Document(file_path)
        return "\n".join(p.text for p in d.paragraphs).strip()
    return ""


# Definition of 5 Diverse Candidates
DIVERSE_CANDIDATES = [
    {
        "full_name": "Trần Đức Anh (David Tran)",
        "email": "david.tran.tech@example.com",
        "phone": "0918112233",
        "skills": "Python, FastAPI, Django, Flask, PostgreSQL, Redis, Docker, Kubernetes, AWS, Kafka, Microservices, CI/CD, PyTest",
        "experience": "5 năm Senior Backend Engineer tại Fintech & SaaS",
        "education": "B.S. in Computer Science - Đại học Công nghệ (ĐHQGHN)",
        "source": "LINKEDIN",
        "format": "pdf",
        "filename": "david_tran_backend_lead.pdf",
        "title": "Senior Backend Engineer",
        "job_id": 1,  # Python Developer
        "sections": [
            (
                "PROFESSIONAL SUMMARY",
                "Results-driven Senior Backend Engineer with 5+ years of experience architecting resilient distributed systems and RESTful APIs using Python, FastAPI, and Flask. Strong track record in fintech systems, payment processing, and high-concurrency database optimization.",
            ),
            (
                "WORK EXPERIENCE",
                "Senior Backend Engineer - PayFlow Vietnam (2022 - Present):\n"
                "- Re-architected core transaction pipeline handling 10M+ daily events using FastAPI, Kafka, and Redis, reducing p99 latency from 320ms to 45ms.\n"
                "- Built automated reconciliation microservice reconciling $15M daily volume with 99.999% accuracy.\n"
                "- Spearheaded CI/CD pipelines with GitHub Actions and Docker, cutting deployment cycle by 40%.\n"
                "Backend Software Engineer - NexaTech (2020 - 2022):\n"
                "- Developed REST APIs using Flask and PostgreSQL for enterprise CRM supporting 200,000 active users.\n"
                "- Optimized slow database queries and indexing strategies, decreasing average response time by 35%.",
            ),
            (
                "TECHNICAL SKILLS",
                "Languages: Python (Advanced), SQL, Bash.\n"
                "Frameworks: FastAPI, Flask, Django, Celery, PyTest.\n"
                "Databases: PostgreSQL, MySQL, Redis, MongoDB.\n"
                "DevOps & Cloud: Docker, Kubernetes, AWS (EC2, S3, RDS), Kafka, GitHub Actions.",
            ),
            (
                "EDUCATION & CERTIFICATIONS",
                "VNU University of Engineering and Technology (UET) - B.S. in Computer Science (GPA: 3.6/4.0).\n"
                "AWS Certified Solutions Architect - Associate.",
            ),
        ],
    },
    {
        "full_name": "Nguyễn Hoàng Mai Phương",
        "email": "maiphuong.marketing@example.com",
        "phone": "0933224455",
        "skills": "Digital Marketing, Meta Ads, TikTok Ads, Content Creator, SEO Onpage, Google Analytics 4, CapCut, Canva, Copywriting",
        "experience": "3.5 năm Quản lý Chiến dịch & Sáng tạo Nội dung Đa nền tảng",
        "education": "Cử nhân Marketing - Đại học Kinh tế Quốc dân (NEU)",
        "source": "FACEBOOK",
        "format": "docx",
        "filename": "mai_phuong_marketing_growth.docx",
        "title": "Senior Digital Marketing & Growth Executive",
        "job_id": 2,  # Marketing Executive
        "sections": [
            (
                "MỤC TIÊU NGHỀ NGHIỆP",
                "Chuyên viên Digital Marketing năng động với hơn 3 năm kinh nghiệm trong lĩnh vực Social Media, Performance Ads và Branding. Đam mê xây dựng nội dung truyền cảm hứng và thúc đẩy tăng trưởng doanh số qua các kênh chuyển đổi số.",
            ),
            (
                "KINH NGHIỆM LÀM VIỆC",
                "Chuyên viên Digital Marketing - GenZ Media & Retail (2023 - Hiện tại):\n"
                "- Trực tiếp quản lý và sản xuất nội dung kênh TikTok thương hiệu, tăng trưởng từ 0 lên 320.000 followers trong 6 tháng, 45 video đạt triệu views.\n"
                "- Vận hành ngân sách quảng cáo Meta & TikTok Ads 150 triệu VNĐ/tháng, duy trì ROAS trung bình 4.2x.\n"
                "- Triển khai chiến dịch ra mắt bộ sưu tập hè 2024, mang về 1.2 tỷ doanh số online trong 2 tuần đầu.\n"
                "Content & Social Executive - BeautyHub Vietnam (2021 - 2023):\n"
                "- Viết bài chuẩn SEO cho website, nâng thứ hạng 50+ từ khóa chiến lược vào Top 3 Google, tăng organic traffic 180%.\n"
                "- Lên kịch bản livestream bán hàng, đạt kỷ lục 15.000 mắt xem trực tiếp.",
            ),
            (
                "KỸ NĂNG CHUYÊN MÔN",
                "Quản lý quảng cáo: Meta Ads Manager, TikTok Ads Manager, Google Ads.\n"
                "Phân tích & Tối ưu: Google Analytics 4 (GA4), Search Console, Ahrefs, A/B Testing.\n"
                "Sáng tạo nội dung: CapCut, Canva, Photoshop cơ bản, Kỹ năng Copywriting.",
            ),
            (
                "HỌC VẤN & BẰNG CẤP",
                "Đại học Kinh tế Quốc dân (NEU) - Tốt nghiệp loại Giỏi chuyên ngành Marketing.\n"
                "Chứng chỉ Google Analytics Certified, Meta Certified Digital Marketing Associate.",
            ),
        ],
    },
    {
        "full_name": "Đỗ Hoàng Yến",
        "email": "hoangyen.hr@example.com",
        "phone": "0977665544",
        "skills": "Tech Sourcing, STAR Interview, Talent Acquisition, Employer Branding, ATS, Onboarding, C&B, Headhunting",
        "experience": "4 năm Chuyên viên Tuyển dụng & Thu hút Nhân tài IT",
        "education": "Cử nhân Quản trị Kinh doanh - Đại học Ngoại thương (FTU)",
        "source": "JOB_SITE",
        "format": "pdf",
        "filename": "hoang_yen_hr_specialist.pdf",
        "title": "Senior HR & Talent Acquisition Specialist",
        "job_id": 3,  # HR Recruiter
        "sections": [
            (
                "GIỚI THIỆU BẢN THÂN",
                "Chuyên viên Tuyển dụng và Thu hút Nhân tài (Talent Acquisition) với 4 năm kinh nghiệm chuyên sâu trong ngành Công nghệ Thông tin (IT Recruitment). Có mạng lưới quan hệ rộng khắp với hơn 5.000 kỹ sư công nghệ trên LinkedIn và các cộng đồng lập trình viên.",
            ),
            (
                "KINH NGHIỆM CÔNG TÁC",
                "Senior IT Recruiter - VinaTech Solutions (2022 - Hiện tại):\n"
                "- Đảm nhận toàn diện quy trình tuyển dụng cho các vị trí Backend (Python, Java), Frontend (React), DevOps và Product Owner.\n"
                "- Tuyển dụng thành công 65 kỹ sư trong năm 2024 (đạt 115% KPI), tỷ lệ ứng viên vượt qua 2 tháng thử việc đạt 94%.\n"
                "- Rút ngắn thời gian tuyển dụng (Time-to-Hire) từ 38 ngày xuống 20 ngày nhờ chuẩn hóa quy trình sàng lọc và xây dựng Talent Pool.\n"
                "- Phối hợp với Tech Lead xây dựng khung đánh giá năng lực và bộ câu hỏi phỏng vấn theo mô hình hành vi STAR.\n"
                "Recruitment Executive - TalentLink HR Agency (2020 - 2022):\n"
                "- Headhunting nhân sự cấp trung và cấp cao cho các khách hàng công nghệ tài chính.\n"
                "- Tổ chức các sự kiện Job Fair tại các trường đại học lớn thu hút 1.200 sinh viên tham gia.",
            ),
            (
                "KỸ NĂNG CỐT LÕI",
                "Sourcing & Headhunting: LinkedIn Recruiter, GitHub Sourcing, Boolean Search.\n"
                "Phỏng vấn & Đánh giá: Phương pháp STAR, Đánh giá văn hóa doanh nghiệp (Culture Fit).\n"
                "Vận hành nhân sự: Quản lý hệ thống ATS, Onboarding bài bản, Đàm phán Offer & C&B.",
            ),
            (
                "HỌC VẤN & CHỨNG CHỈ",
                "Đại học Ngoại thương (FTU) - Chuyên ngành Quản trị Kinh doanh Quốc tế.\n"
                "Chứng chỉ Quản trị Nhân sự Chuyên nghiệp (CPHR), Khóa học Phỏng vấn Hành vi Nâng cao.",
            ),
        ],
    },
    {
        "full_name": "Lâm Quốc Bảo",
        "email": "quocbao.sales@example.com",
        "phone": "0988776655",
        "skills": "B2B Solution Sales, Key Account Management, Contract Negotiation, CRM Salesforce, Lead Generation, Pipeline Management",
        "experience": "6 năm Giám đốc Kinh doanh B2B & Phát triển Thị trường",
        "education": "Cử nhân Thương mại (Commerce) - Đại học RMIT Việt Nam",
        "source": "REFERRAL",
        "format": "docx",
        "filename": "quoc_bao_b2b_sales.docx",
        "title": "Senior B2B Sales & Business Development Manager",
        "job_id": 2,  # Marketing Executive / Commercial
        "sections": [
            (
                "TÓM TẮT NĂNG LỰC",
                "Chuyên gia Kinh doanh B2B giàu năng lượng và kỷ luật với hơn 6 năm thành công trong lĩnh vực bán giải pháp phần mềm doanh nghiệp (B2B SaaS / ERP / CRM). Đam mê mở rộng thị trường, xây dựng quan hệ đối tác chiến lược và dẫn dắt đội ngũ bán hàng đạt doanh số đột phá.",
            ),
            (
                "THÀNH TÍCH & KINH NGHIỆM",
                "B2B Sales Manager - CloudEnterprise Vietnam (2021 - Hiện tại):\n"
                "- Đạt doanh số cá nhân 8.5 tỷ VNĐ năm 2024, vượt 135% chỉ tiêu năm.\n"
                "- Chốt thành công 25 hợp đồng doanh nghiệp lớn (Enterprise Accounts) trong các ngành Bán lẻ, Logistics và Sản xuất.\n"
                "- Duy trì tỷ lệ tái ký hợp đồng (Retention Rate) đạt 92%, nâng cao giá trị vòng đời khách hàng (LTV) thêm 30%.\n"
                "- Quản lý và đào tạo đội ngũ 5 chuyên viên Account Executive, liên tục đạt top phòng kinh doanh xuất sắc nhất quý.\n"
                "Senior Business Development Executive - FastERP (2018 - 2021):\n"
                "- Tìm kiếm khách hàng tiềm năng qua cold outreach, LinkedIn B2B và mạng lưới đối tác.\n"
                "- Đàm phán trực tiếp với Giám đốc điều hành (CEO) và Giám đốc Công nghệ (CTO) các tập đoàn vừa và lớn.",
            ),
            (
                "KỸ NĂNG CHUYÊN MÔN",
                "Kỹ năng thương mại: Đàm phán hợp đồng lớn, Quản lý phễu bán hàng (Pipeline Management), Thuyết trình giải pháp bán hàng (Pitching).\n"
                "Công cụ & CRM: Salesforce, HubSpot CRM, LinkedIn Sales Navigator, Asana.",
            ),
            (
                "HỌC VẤN",
                "Đại học RMIT Việt Nam - Cử nhân Thương mại (Bachelor of Commerce).\n"
                "Chứng chỉ Solution Selling & High-Stake Negotiation.",
            ),
        ],
    },
    {
        "full_name": "Vũ Hoàng Long",
        "email": "hoanglong.data@example.com",
        "phone": "0903344556",
        "skills": "Python, SQL, Power BI, Financial Modeling, CFA Level 1, Pandas, Data Automation, Quantitative Analysis",
        "experience": "3 năm Chuyên viên Phân tích Dữ liệu Tài chính & Tự động hóa",
        "education": "Thạc sĩ Tài chính Ứng dụng - Đại học Kinh tế TP.HCM (UEH)",
        "source": "WEBSITE",
        "format": "pdf",
        "filename": "hoang_long_data_finance.pdf",
        "title": "Financial & Data Analytics Specialist",
        "job_id": 7,  # Backend / Data
        "sections": [
            (
                "PROFESSIONAL SUMMARY",
                "Data & Financial Analyst with 3 years of rigorous experience blending financial modeling with advanced data automation using Python (Pandas) and SQL. Passionate about transforming complex transaction datasets into actionable strategic intelligence for executive leadership.",
            ),
            (
                "WORK EXPERIENCE",
                "Chuyên viên Phân tích Dữ liệu Tài chính - FinaCorp Vietnam (2022 - Hiện tại):\n"
                "- Thiết kế và triển khai hệ thống 12 Dashboard Power BI tự động kết nối cơ sở dữ liệu SQL, phục vụ báo cáo doanh thu và chi phí thời gian thực cho Ban Giám đốc.\n"
                "- Giảm thiểu 75% thời gian xử lý dữ liệu thủ công mỗi tuần nhờ viết các script Python tự động hóa trích xuất và đối soát số liệu.\n"
                "- Tham gia xây dựng mô hình dự báo dòng tiền và quản trị rủi ro thanh khoản, nâng cao độ chính xác dự báo lên 92%.\n"
                "Financial Analyst - SmartInvest Capital (2021 - 2022):\n"
                "- Thẩm định tài chính và định giá doanh nghiệp (DCF, Multiples) cho 15 thương vụ đầu tư tiềm năng.\n"
                "- Viết báo cáo phân tích ngành và dự báo xu hướng thị trường tài chính định kỳ.",
            ),
            (
                "CHỨNG CHỈ & HỌC VẤN",
                "Passed CFA Level 1 (Chartered Financial Analyst).\n"
                "IELTS 7.5 Academic.\n"
                "Thạc sĩ Tài chính Ứng dụng - Đại học Kinh tế TP.HCM (UEH), Tốt nghiệp loại Giỏi.",
            ),
            (
                "KỸ NĂNG PHÂN TÍCH & CÔNG NGHỆ",
                "Lập trình & Cơ sở dữ liệu: Python (Pandas, NumPy, Matplotlib), SQL nâng cao (Window Functions, CTE, Indexing).\n"
                "Trực quan hóa & Báo cáo: Microsoft Power BI, DAX, Excel Financial Modeling, VBA.",
            ),
        ],
    },
]


def seed():
    print("=== BẮT ĐẦU TẠO CÁC BẢN CV ĐA DẠNG MẪU ===")
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)

    for cand_info in DIVERSE_CANDIDATES:
        fname = cand_info["filename"]
        fmt = cand_info["format"]

        # 1. Generate file
        if fmt == "pdf":
            file_path = create_pdf_cv(
                fname,
                cand_info["full_name"],
                cand_info["title"],
                cand_info["email"],
                cand_info["phone"],
                cand_info["sections"],
            )
        else:
            file_path = create_docx_cv(
                fname,
                cand_info["full_name"],
                cand_info["title"],
                cand_info["email"],
                cand_info["phone"],
                cand_info["sections"],
            )

        print(f"[+] Đã tạo file CV: {fname} ({fmt.upper()})")

        # 2. Extract CV text
        extracted_text = extract_text(file_path)
        print(f"    -> Đã trích xuất {len(extracted_text)} ký tự văn bản")

        # 3. Insert or update candidate in MySQL
        cv_rel_path = f"uploads/{fname}"
        cursor.execute("SELECT id FROM candidates WHERE email = %s", (cand_info["email"],))
        row = cursor.fetchone()

        if row:
            cid = row["id"]
            cursor.execute(
                """
                UPDATE candidates
                SET full_name = %s, phone = %s, skills = %s, experience = %s, education = %s,
                    source = %s, cv_file = %s, cv_text = %s
                WHERE id = %s
                """,
                (
                    cand_info["full_name"],
                    cand_info["phone"],
                    cand_info["skills"],
                    cand_info["experience"],
                    cand_info["education"],
                    cand_info["source"],
                    cv_rel_path,
                    extracted_text,
                    cid,
                ),
            )
            print(f"    -> Cập nhật ứng viên ID: {cid} ({cand_info['full_name']})")
        else:
            cid = create_candidate(
                full_name=cand_info["full_name"],
                email=cand_info["email"],
                phone=cand_info["phone"],
                skills=cand_info["skills"],
                experience=cand_info["experience"],
                education=cand_info["education"],
                source=cand_info["source"],
                cv_file=cv_rel_path,
                cv_text=extracted_text,
            )
            print(f"    -> Tạo mới ứng viên ID: {cid} ({cand_info['full_name']})")

        conn.commit()

        # 4. Check & create application if not exists
        jid = cand_info["job_id"]
        if not application_exists(cid, jid):
            app_id = create_application(
                candidate_id=cid,
                job_id=jid,
                note=f"Hồ sơ ứng tuyển mẫu ({fmt.upper()}) - Nguồn {cand_info['source']}",
            )
            print(f"    -> Tạo đơn ứng tuyển ID: {app_id} cho vị trí Job ID: {jid}")
        else:
            print(f"    -> Đã có đơn ứng tuyển cho vị trí Job ID: {jid}")

    cursor.close()
    conn.close()

    # 5. Rebuild RAG Vector index
    print("\n=== ĐỒNG BỘ LẠI CHỈ MỤC TÌM KIẾM VECTOR RAG (FAISS) ===")
    rag_info = rebuild_index()
    print(f"Vector RAG index đã cập nhật: {rag_info.get('documents')} documents, {rag_info.get('entity_counts', {}).get('cv_chunks')} CV chunks.")
    print("=== HOÀN TẤT ===")


if __name__ == "__main__":
    seed()
